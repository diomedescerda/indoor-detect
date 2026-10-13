"""Build the paper handoff report from the archived benchmark tables."""
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT/'benchmark/results/android'
CONFIGS = [(m,l) for m in ('head_harness','chest_harness','gorillapod') for l in ('1x','0.5x')]
NAMES = ['Head 1x','Head 0.5x','Chest 1x','Chest 0.5x','Gorillapod 1x','Gorillapod 0.5x']


def read(name):
    return list(csv.DictReader((RESULTS/name).open()))


def main():
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.7)
    sec.left_margin=sec.right_margin=Inches(.72)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        style=doc.styles[name];style.font.name='Liberation Sans';style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(25)
    doc.styles['Caption'].font.size=Pt(9.5)
    doc.styles['Caption'].font.bold=False
    doc.styles['Caption'].font.italic=False
    doc.styles['Heading 1'].font.size=Pt(15)
    doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run('AssistiveVision  |  ');run.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    doc.core_properties.title='AssistiveVision capture configuration study'
    doc.core_properties.subject='Android benchmark and paper handoff'
    doc.core_properties.author='AssistiveVision research project'

    def para(text,style=None):return doc.add_paragraph(text,style)
    def heading(text):doc.add_heading(text,1)
    def table(headers,rows,widths,size=9.5):
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Inches(w)
        for c,h in zip(t.rows[0].cells,headers):c.text=h
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        borders=OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        t._tbl.tblPr.append(borders)
        for ri,row in enumerate(t.rows):
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for ci,cell in enumerate(row.cells):
                cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd')
                shade.set(qn('w:fill'),'DDE8EE' if ri==0 else ('F5F7F8' if ri%2==0 else 'FFFFFF'));tc.append(shade)
                margins=OxmlElement('w:tcMar')
                for edge,val in [('top',65),('bottom',65),('left',85),('right',85)]:
                    node=OxmlElement('w:'+edge);node.set(qn('w:w'),str(val));node.set(qn('w:type'),'dxa');margins.append(node)
                tc.append(margins)
                for p in cell.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
                    for r in p.runs:r.font.size=Pt(size);r.bold=ri==0;r.font.color.rgb=RGBColor(0,0,0)
        para('').paragraph_format.space_after=Pt(0)
        return t

    summary=read('summary.csv')
    sm={(r['mode'],r['mount'],r['lens']):r for r in summary if r['aggregation']=='configuration'}
    scenes=read('scenes/scene_counts.csv')
    lookup={(r['mode'],int(r['capture_group']),r['mount'],r['lens']):r for r in scenes}
    pairs=[r for r in read('scenes/paired_count_differences.csv') if r['mode']=='detector' and r['metric']=='tp']
    pct=lambda v:f'{float(v)*100:.1f}%'

    doc.add_heading('AssistiveVision capture configuration study',0)
    para('Benchmark and paper handoff   13 October 2026','Subtitle')
    para('Gorillapod with the 0.5x lens is the leading candidate when the goal is to detect more objects across scenes. Chest harness with the 1x lens has the highest F1. This report preserves both findings, the paired scene evidence, and the method needed to continue the paper from the repository.')
    heading('Study and primary results')
    para('The study contains 198 iPhone photos in 33 successive six-photo capture groups. Each group compares head harness, chest harness and gorillapod, each at 1x and 0.5x. The corrected ground truth contains 889 instances in ten classes, including 16 images with no annotated target objects. All 198 photos completed Android replay on a Pixel 8 emulator.')
    data=[]
    for c,name in zip(CONFIGS,NAMES):
        r=sm['detector',*c];a=sm['app_output',*c]
        data.append([name,r['gt'],r['tp'],r['fp'],pct(r['f1']),a['tp']])
    table(['Configuration','Visible','Correct','False','F1','Correct\ntop two'],data,[1.8,.83,.83,.73,.82,1.0])
    para('Visible means reviewed instances inside each photo. Correct and false mean matched and unmatched detections at confidence 0.30 and IoU 0.50. Counts are totals over 33 photos per configuration. The final column applies the app default limit of two selected objects; it does not measure spoken announcements.','Caption')
    heading('Decision based on object counts')
    para('Gorillapod 0.5x returns 82 correct detections, compared with 65 for chest 1x and 77 for chest 0.5x. It produces at least one correct detection in 29 of 33 scenes, and reaches the best nonzero count in 20 scenes, including ties. Its 174 visible annotations represent the largest observed coverage across configurations.')
    para('Against chest 1x, the gain is 17 correct detections, or 26.2%. The paired mean difference is 0.515 objects per scene, with an exploratory 95% bootstrap interval of 0.121 to 0.939. Against chest 0.5x, the gain is five detections, with a mean difference of 0.152 and an interval of -0.182 to 0.515. That smaller advantage remains uncertain.')
    para('The F1 leader, chest 1x, achieves precision 75.6%, recall 52.4%, F1 61.9% and COCO AP 42.2%. Gorillapod 0.5x achieves precision 72.6%, recall 47.1%, F1 57.1% and AP 41.5%. The preferred setup therefore depends on whether the priority is count or proportional detection performance.')

    doc.add_page_break()
    heading('Paired scene evidence')
    para('Each row below compares Gorillapod 0.5x with one alternative in the same 33 acquisition groups. More, equal and fewer refer to correctly detected object counts. Confidence intervals resample whole groups together, preserving the pairing.')
    names=dict(zip(CONFIGS,NAMES));data=[]
    for r in pairs:
        m,l=r['configuration_b'].split('/')
        data.append([names[m,l],f"{r['scenes_a_greater']} / {r['scenes_tied']} / {r['scenes_a_lower']}",f"{int(r['total_difference']):+d}",f"{float(r['mean_difference']):+.3f}",f"{float(r['ci95_low']):+.3f} to {float(r['ci95_high']):+.3f}"])
    table(['Alternative','More / equal / fewer','Total gain','Gain per scene','95% interval'],data,[1.52,1.35,.8,1.04,1.3])
    para('These are exploratory percentile intervals from 10,000 paired bootstrap samples, seed 20261013. They are unadjusted for multiple comparisons and assume acquisition groups can be resampled independently. Groups that share a room may violate that assumption.','Caption')
    heading('Examples and counterexamples')
    data=[]
    for s,label in [(3,'03 Living room'),(10,'10 Dining room'),(12,'12 Living room'),(28,'28 Outdoor seating'),(30,'30 Salon')]:
        vals=[lookup['detector',s,*c] for c in [('chest_harness','1x'),('chest_harness','0.5x'),('gorillapod','0.5x')]]
        data.append([label]+[f"{r['tp']} / {r['gt']}" for r in vals])
    table(['Scene','Chest 1x','Chest 0.5x','Gorillapod 0.5x'],data,[2.0,1.25,1.25,1.51])
    para('Cells show correct detections / visible annotations. Scene names are descriptive labels from inspected photos, not an independently recorded scene taxonomy.','Caption')
    para('In scene 10, both chest views contain objects but yield zero correct detections. Gorillapod 0.5x detects four, so this difference involves recognition of visible objects. In scene 12, it detects nine of ten annotations, compared with six of eight and seven of ten for the chest views.')
    para('Scene 28 shows the limit of wider coverage: Gorillapod 0.5x contains 36 annotations versus 23 for chest 1x, but correct detections rise only from nine to ten. Small or distant objects remain difficult. In scene 30, chest 0.5x performs better, with seven correct and one false detection versus four correct and six false detections for Gorillapod 0.5x.')
    para('With the top-two cap, Gorillapod 0.5x has 49 correct and five false selections; chest 0.5x has 44 and eight; chest 1x has 37 and seven. Across all configurations, the cap lowers recall from 47.8% to 26.8% while raising precision from 71.5% to 84.7%. Announcing fewer objects may still serve usability goals that this benchmark does not test.')

    doc.add_page_break()
    heading('Dataset and evaluation method')
    para('Photos retain their original 1536 by 2048 resolution and require no EXIF rotation. Ascending IMG numbers are grouped in the supplied order: head 1x, head 0.5x, chest 1x, chest 0.5x, gorillapod 1x, gorillapod 0.5x. The group assignments should be checked against capture notes before publication.')
    para('Annotations began as a model-assisted draft in Codex and were corrected by the user in Roboflow. The canonical COCO export maps categories by name to person 1, chair 62, couch 63, bed 65, dining table 67, toilet 70, tv 72, laptop 73, sink 81 and refrigerator 82. Thirteen boxes affected by tiny export rounding were clipped to image bounds. There is no independent second-reviewer audit. The exact serving chat model is not independently established by the saved files.')
    para('The Android instrumentation test calls the production preprocessing, detection and inverse letterbox modules. EfficientDet-Lite2 receives a 448 by 448 uint8 RGB tensor; portrait images are letterboxed to 336 by 448 with gray 114 padding. Three warm-up passes precede inference, and image order is shuffled with seed 20261013. Low-score predictions are retained at 0.001, with the production maximum of 15 target-class detections.')
    para('Operating metrics filter predictions at confidence 0.30 and match them greedily in score order to unmatched annotations of the same class at IoU at least 0.50. Precision is TP/(TP+FP), recall is TP/(TP+FN), and F1 is 2TP/(2TP+FP+FN). COCO AP averages IoU thresholds 0.50 through 0.95 using the saved low-score predictions; AP50 and AP75 are also archived. Empty images remain in evaluation.')
    heading('Interpretation limits')
    para('The object-count objective was emphasized after the initial F1 comparison. Both analyses are exploratory. Choosing a setup on this dataset makes it a configuration-selection set; confirm the choice on new scenes. Chairs account for 438 of 889 instances, so pooled counts and F1 give frequent classes more weight. Per-class results accompany the totals.')
    para('Annotations count instances visible within each photograph. Objects outside a frame cannot count as detector false negatives. Shared object identities across the six views were not annotated, so the data cannot measure recovery of the same unique scene objects. Successive photos also allow people and poses to change. Observed count differences cannot be attributed solely to mount or lens.')
    para('Pixel_8 is an emulator profile running Android 17, API 37, x86_64 with 16 KB pages. The inputs remain iPhone photographs; Pixel camera optics are not simulated. Detection-method median latency was 246.3 ms and p95 307.5 ms; median JPEG decoding plus preprocessing was 104.0 ms. These debug emulator timings exclude camera capture, UI and speech and do not estimate physical Pixel 8 speed.')
    heading('Repository handoff')
    para('Build/Testing contains the original photos, final labels, original reviewed export, historical drafts, predictions, metric tables, figures, all 33 comparison sheets, model asset, Android test and analysis scripts. README.md describes how to reproduce the analysis on another computer. SHA256SUMS.json and verify_package.py check package integrity. The primary reporting source is benchmark/results/android.json; desktop results are a separate proxy comparison.')
    para('The tested Android app commit is 85c852ddcb543bce1bd69ad0db307f895d597881 in diomedescerda/AssistiveVision. App source snapshots and hashes are preserved in provenance/. Instrumentation completed successfully for all 198 images; four evaluation tests passed. The interrupted run is excluded. Document annotation conventions and upstream model provenance when preparing the paper.')

    doc.add_page_break()
    heading('Correct detections in all 33 scenes')
    para('Each cell is the number of correct detections at confidence 0.30 and IoU 0.50 before the app top-two limit. Ground-truth, false-positive and missed-object counts are available in benchmark/results/android/scenes/scene_counts.csv. All scene sheets use the same numbering.')
    data=[]
    for s in range(1,34):data.append([f'{s:02d}']+[lookup['detector',s,*c]['tp'] for c in CONFIGS])
    table(['Scene','Head\n1x','Head\n0.5x','Chest\n1x','Chest\n0.5x','Gorillapod\n1x','Gorillapod\n0.5x'],data,[.57,.84,.89,.84,.89,1.0,.98],size=9)
    para('Gorillapod 0.5x is the provisional choice for object coverage. Its lead over chest 0.5x and gorillapod 1x remains uncertain, and scene 30 demonstrates a meaningful failure case. A next study should track shared objects across views and validate the selected configuration with the physical target phone.')
    for root in (doc.styles.element,doc.element):
        for border in root.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    dest=ROOT/'AssistiveVision_Benchmark_Report.docx';doc.save(dest);print(dest)


if __name__=='__main__':main()
