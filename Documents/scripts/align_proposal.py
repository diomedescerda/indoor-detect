"""Add proposal traceability and archived-data analyses to the retained-template build."""
import json
from pathlib import Path
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[2]

def apply(doc, original, add_before, change, H2, H3, CAP):
    evidence=json.loads((ROOT/'Documents/Review/proposal_evidence/summary.json').read_text())
    def table(anchor, caption, headers, rows):
        add_before(anchor,caption,CAP)
        t=doc.add_table(rows=1,cols=len(headers))
        for c,v in zip(t.rows[0].cells,headers):c.text=v
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        anchor.addprevious(t._tbl)
        return t
    def fmt(value):return f'{value:.1f}'.replace('.',',')
    change(10,'Diomedes Farid Cerda Leon');change(27,'Diomedes Farid Cerda Leon')
    change(35,'Idanis Beatriz Diaz Bolaño');change(37,'Aquiles Alfonso Cohen Llanes')
    change(43,'Desarrollo y gestión de tecnologías para las organizaciones y la sociedad – TecnOS')
    change(258,'La metodología conserva las tres fases de la propuesta de pasantía [34]. La fase 1 comprende revisión de literatura, análisis del problema, requerimientos, selección del modelo, delimitación de objetos y diseño de descriptores. La comparación inicial de seis configuraciones sobre 80 imágenes y 257 anotaciones pertenece a esta fase y se presenta en el capítulo 2. La fase 2 corresponde al diseño integrado y al desarrollo del prototipo Android, descritos en el capítulo 3. La fase 3 evalúa el posicionamiento y el desempeño con 198 fotografías organizadas en 33 grupos; sus resultados se presentan en el capítulo 4. El benchmark inicial en Redmi Note 13 y la reproducción posterior en emulador Pixel 8 tienen alcances distintos y se analizan por separado.')
    add_before(original[259]._p,'La correspondencia entre actividades, evidencias y límites se sintetiza en la sección 4.9. La evaluación de la fase 3 cubre detección y procesamiento de imágenes; la iluminación se explora mediante brillo digital, sin mediciones de iluminancia, y la oclusión no cuenta con etiquetas específicas. La claridad auditiva, el tiempo hasta escuchar el mensaje y el efecto sobre la autonomía requieren evaluación adicional. El análisis geométrico de distancia es complementario y no sustituye estas actividades.')
    add_before(original[387]._p,'La prioridad de las diez clases de la Tabla 3.1 responde a una delimitación técnica para explorar mobiliario, equipamiento doméstico y presencia de personas, junto con la disponibilidad de etiquetas COCO [13]. No procede de un estudio de necesidades con usuarios ni establece una jerarquía clínica o de riesgo. La propuesta menciona puertas, sillas y escritorios como ejemplos [34]; la clase silla está incluida, pero el vocabulario COCO empleado no contiene una clase puerta ni una clase independiente escritorio. La etiqueta dining table corresponde a mesa de comedor y no permite afirmar detección general de escritorios. Open Images se consideró como alternativa de vocabulario [20], pero no se entrenó ni evaluó aquí un detector con ese conjunto. Incorporar puertas exige seleccionar o adaptar un modelo compatible, anotar ejemplos y evaluar su desempeño.')
    add_before(original[390]._p,'El diseño de los descriptores combina presencia, nombre del objeto y región relativa al punto de orientación de la cámara. La redacción localizada y el límite de dos objetos por defecto buscan evitar mensajes excesivamente extensos; esa intención de diseño no demuestra una reducción de carga cognitiva. Las pruebas de SpeechFormatter verifican la composición textual, mientras la evaluación de hasta dos objetos cuantifica la selección disponible para el módulo de voz. La comprensión de izquierda, centro, derecha y de las regiones verticales debe comprobarse con oyentes; no se anuncian distancias métricas ni instrucciones de ruta libre.')
    a=original[404]._p
    add_before(a,'3.13 Requerimientos y criterios de verificación',H2)
    add_before(a,'La Tabla 3.4 explicita los requerimientos operativos que se desprenden de la propuesta y de las decisiones implementadas. Es una sistematización de la especificación actual; no representa una sesión de levantamiento de necesidades con usuarios. Se distingue la verificación técnica disponible de los criterios de aceptación que aún requieren ensayos.')
    table(a,'Tabla 3.4. Requerimientos operativos y evidencia de verificación.',['Requerimiento','Criterio verificable','Evidencia y alcance'],[
      ('Captura bajo demanda','Procesar una captura y liberar la vista al terminar o fallar','CameraX y flujo manual, secciones 3.3 y 3.11; falta prueba de interacción con usuarios.'),
      ('Detección local','Ejecutar el modelo seleccionado y conservar cajas, clases y puntuaciones','EfficientDet-Lite2 y reproducción de 198 imágenes, sección 3.5 y capítulo 4.'),
      ('Ubicación relativa','Asignar una de seis zonas al centro de la caja','LocalizationModule y seis pruebas JVM; no demuestra precisión de ubicación física.'),
      ('Descripción auditiva','Construir mensajes localizados y sintetizar voz en el dispositivo','SpeechFormatter, AudioModule y catálogo de voces; claridad auditiva pendiente.'),
      ('Salida configurable','Aplicar confianza y máximo de objetos elegidos','Lista de uno a tres objetos, dos por defecto; análisis de selección en sección 4.5.'),
      ('Rendimiento y montaje','Medir tiempos con alcance explícito y comparar seis configuraciones','Capítulos 2 y 4; faltan latencia completa, ergonomía y estabilidad en uso.'),
    ])
    add_before(a,'Los objetivos no establecieron un umbral numérico de aceptación para claridad auditiva o tiempo total de respuesta. No se introduce uno después de observar los resultados. El siguiente ensayo debe fijar el umbral, el dispositivo, las condiciones de captura y el protocolo de usuarios antes de evaluar el cumplimiento.')
    a=original[430]._p
    add_before(a,'4.7 Distribución del tiempo de procesamiento',H2)
    add_before(a,'Para ampliar la evaluación de tiempo de respuesta prevista en la fase 3, se analizaron los tiempos archivados de las 198 imágenes, sin repetir la inferencia. El benchmark conservó candidatos desde confianza 0,001 para la evaluación; los tiempos no proceden de una nueva ejecución con corte 0,30. La prueba de instrumentación realizó tres ejecuciones de calentamiento y recorrió las fotografías en un orden aleatorio fijo. La Tabla 4.8 presenta mediana, cuartiles y percentil 95; los cuantiles se obtienen por interpolación lineal en la posición (n − 1)p.')
    tt=evidence['timing_ms']
    table(a,'Tabla 4.8. Distribución de tiempos de reproducción en emulador en milisegundos.',['Etapa','P25','Mediana','P75','P95'],[[label]+[fmt(tt[key][k]) for k in ['p25','median','p75','p95']] for label,key in [('Decodificación y preprocesamiento','preprocess_ms'),('Método de detección','inference_ms'),('Suma por imagen','replay_total_ms')]])
    add_before(a,'La suma se calculó para cada imagen antes de resumirla: su mediana es '+fmt(tt['replay_total_ms']['median'])+' ms y su P95 es '+fmt(tt['replay_total_ms']['p95'])+' ms. No se sumaron las medianas de etapas. El intervalo instrumentado de detección incluye la llamada y preparación inmediata de su medición; no es un perfil de operadores del modelo. Estos tiempos corresponden a JPEG reproducidos en un emulador Pixel 8 con compilación de depuración. Excluyen adquisición de cámara, interfaz, localización, generación de voz y reproducción del mensaje. Por tanto, no representan el tiempo desde pulsar Capturar hasta oír la respuesta, ni el rendimiento de un Pixel 8 físico.')
    add_before(a,'4.8 Brillo de las imágenes y alcance del análisis de oclusión',H2)
    add_before(a,'La propuesta incluye iluminación y oclusión como condiciones de evaluación [34]. Las fotografías no tienen registros de lux ni etiquetas de oclusión. Como análisis posterior disponible, se calculó el brillo digital medio de cada JPEG a resolución original mediante la conversión RGB a escala de grises de ocho bits de Pillow [35]. Se dividió la muestra en tres grupos mediante los cuantiles 1/3 y 2/3 del brillo medio. Son grupos relativos a esta muestra; no se denominan iluminación baja, normal o alta, porque el brillo depende también de exposición, superficies y composición de la escena.')
    table(a,'Tabla 4.9. Detección agrupada por brillo digital de las fotografías.',['Grupo de brillo','Fotos','GT','TP','FP','Recall'],[[r['group'],r['n'],r['gt'],r['tp'],r['fp'],fmt(100*r['recall'])+' %'] for r in evidence['brightness_groups']])
    b=evidence['brightness_groups']
    add_before(a,'Cada grupo reúne 66 fotografías. Los intervalos observados de brillo medio son '+ '; '.join(f"{r['group'].lower()}: {fmt(r['luma_min'])} a {fmt(r['luma_max'])}" for r in b)+' en una escala de 0 a 255. La Tabla 4.9 utiliza detecciones con confianza 0,30 e IoU mínimo 0,50 y calcula recall agrupando instancias, sin promediar recalls por fotografía. Se conservan las imágenes vacías en el conteo de fotos; no aportan instancias al denominador de recall.')
    chairs=[r['class_counts']['chair']/r['gt']*100 for r in b]
    add_before(a,'El grupo intermedio presenta mayor recall observado, pero no se obtiene una relación monótona con el brillo. La mezcla de clases también cambia: las sillas representan '+', '.join(fmt(v)+' %' for v in chairs)+' de las instancias de los tres grupos, respectivamente. El montaje, lente, fondo y densidad de objetos varían junto con el brillo. La comparación es descriptiva, no una prueba causal de iluminación ni un umbral operativo; no permite afirmar tolerancia a oscuridad o condiciones de lux específicas. El archivo complementario registra los conteos por clase y configuración para hacer visible esta confusión.')
    add_before(a,'La oclusión requiere una revisión adicional de las imágenes originales. Las cajas COCO disponibles no incluyen fracción visible, límites completos de objetos ocultos ni etiquetas de severidad. El área de una caja pequeña puede deberse a distancia, tamaño físico u oclusión, por lo que no se usa como sustituto de oclusión. Con las fotografías actuales se puede crear una anotación manual de ausencia, oclusión parcial y truncamiento por borde, con definiciones previas y revisión independiente, y después estratificar detección por clase y tamaño aparente. Esa anotación no se presenta como realizada. Las pruebas con iluminación controlada y distancias medidas requieren capturas nuevas.')
    add_before(a,'4.9 Correspondencia con las actividades de la propuesta',H2)
    add_before(a,'La Tabla 4.10 relaciona las catorce actividades de las tres fases con los resultados del informe. «Documentada» indica evidencia del análisis o diseño; «implementada» indica disponibilidad técnica; «parcial» identifica componentes evaluados junto con criterios todavía pendientes. Ninguno de estos estados equivale a demostrar autonomía, seguridad de navegación o aceptación por usuarios.')
    table(a,'Tabla 4.10. Evidencias y alcance de las actividades de la pasantía.',['Fase y actividad','Evidencia','Estado'],[
      ('1. Revisión teórica y problema','Capítulo 1 y antecedentes de la introducción.','Documentada'),
      ('1. Requerimientos y diseño conceptual','Arquitectura, alcance y Tabla 3.4.','Documentada; sin consulta a usuarios'),
      ('1. Selección y evaluación DNN','Seis configuraciones, 80 imágenes, benchmark en Redmi; capítulo 2.','Evaluada en el entorno descrito'),
      ('1. Objetos prioritarios','Diez clases COCO, Tabla 3.1 y sección 3.7.','Parcial; sin puertas ni validación de prioridades'),
      ('1. Diseño de descriptores','Seis zonas y composición textual, secciones 3.6 y 3.8.','Documentada; comprensión pendiente'),
      ('2. Sistema integrado y posición','Flujo modular y seis configuraciones, secciones 3.1 y 3.2.','Implementada; ergonomía pendiente'),
      ('2. Interfaz de captura','CameraX, captura manual y estados, secciones 3.3 y 3.11.','Implementada'),
      ('2. Adaptación e integración DNN','Modelo TFLite, letterbox y coordenadas, secciones 3.4 y 3.5.','Implementada; sin entrenamiento nuevo'),
      ('2. Módulo de audio','Síntesis local y formato de mensajes, sección 3.8.','Implementada; claridad no evaluada'),
      ('2. Prototipo móvil','Integración Android y pruebas, sección 3.12 y anexo B.','Implementada; no validada con usuarios'),
      ('3. Posicionamiento del dispositivo','33 grupos y seis configuraciones, secciones 4.2 y 4.3.','Parcial; sin estabilidad ni comodidad'),
      ('3. Escenarios según objetos','198 fotografías, 889 instancias y diez clases, secciones 4.1 y 4.5.','Documentada; predominan sillas'),
      ('3. Oclusión, iluminación y respuesta','Métricas, tiempos y brillo, secciones 4.5, 4.7 y 4.8.','Parcial; faltan oclusión, lux y latencia con voz'),
      ('3. Informe y recomendaciones','Resultados, límites, conclusiones y trabajo futuro.','Documentada; cierre institucional pendiente'),
    ])
    add_before(a,'El primer objetivo específico cuenta con comparación experimental y selección explícita; el segundo, con un prototipo integrado y verificación técnica. El tercero dispone de resultados de detección y tiempos de procesamiento parciales, pero no de claridad auditiva ni de ensayos controlados de oclusión e iluminación. En consecuencia, el informe sustenta desarrollo y evaluación técnica parcial; el cierre de la evaluación con usuarios debe mantenerse como una actividad pendiente, en correspondencia con el objetivo original.')
    add_before(original[437]._p,'Respecto a las fases propuestas, se documentaron la revisión, la selección del modelo y el desarrollo integrado. La evaluación permanece parcial para claridad auditiva, oclusión, iluminación controlada y tiempo completo de respuesta. El análisis adicional de tiempos obtuvo una mediana de 352,4 ms y un P95 de 419,9 ms para decodificación, preprocesamiento y detección en emulador; el brillo digital permitió describir heterogeneidad sin atribuirla a iluminancia física. Estas ampliaciones aprovechan los datos archivados y delimitan qué evidencia debe obtener el siguiente experimento.')
    add_before(original[441]._p,'Como aprendizaje metodológico, conservar anotaciones, predicciones, temporizaciones y versiones permitió ampliar el informe sin alterar la ejecución original. La ausencia de un protocolo previo de oclusión, lux y escucha limita el cumplimiento de la fase 3. Para próximos proyectos del programa, conviene definir desde el inicio los criterios de aceptación y los instrumentos de medición, separar pruebas técnicas de evaluación con usuarios y reservar un conjunto de escenas independiente para confirmar la configuración seleccionada.')
    doc.add_paragraph('[34] D. F. Cerda Leon, “Sistema de identificación y localización espacial de objetos basado en visión por computador para asistencia a personas con discapacidad visual,” Propuesta de pasantía en investigación, Universidad del Magdalena, documento del proyecto. Disponible en Documents/Propuesta Pasantia de Investigación - Diomedes Cerda.pdf.')
    doc.add_paragraph('[35] Pillow contributors, “ImageStat module,” Pillow documentation. [Online]. Available: https://pillow.readthedocs.io/en/stable/reference/ImageStat.html. El análisis utiliza conversión RGB a L y media de píxeles.')

    # Mark foreign technical terms in the Spanish prose without altering fields or drawings.
    import re
    from copy import deepcopy
    foreign=re.compile(r'\b(?:bounding boxes|bounding box|dining table|recall|bootstrap|letterbox|one-stage|two-stage|on-device|cooldown|frameworks?|pipeline|benchmarks?|backbone|edge|overhead|kernels|datasets?)\b',re.I)
    spanish_body=False
    for p in doc.paragraphs:
        if p.text.strip()=='Introducción':spanish_body=True
        if p.text.strip()=='Bibliografía':spanish_body=False
        if not spanish_body:continue
        for run in list(p.runs):
            matches=list(foreign.finditer(run.text))
            if not matches:continue
            text=run.text;pos=0
            for m in matches:
                for value,italic in [(text[pos:m.start()],False),(m.group(),True)]:
                    if not value:continue
                    clone=deepcopy(run._r)
                    for child in list(clone):
                        if child.tag!=qn('w:rPr'):clone.remove(child)
                    run._r.addprevious(clone)
                    from docx.text.run import Run
                    rr=Run(clone,p);rr.text=value
                    if italic:rr.italic=True
                pos=m.end()
            if pos<len(text):
                clone=deepcopy(run._r);run._r.addprevious(clone);Run(clone,p).text=text[pos:]
            run._r.getparent().remove(run._r)
    # Apply the explicit written rules of the supplied institutional guide.
    for sid, before, after in [('Ttulo1',100,24),('Ttulo2',20,10),('Ttulo3',20,10)]:
        st=doc.styles[sid]
        st.font.name='Arial'
        st.paragraph_format.space_before=Pt(before)
        st.paragraph_format.space_after=Pt(after)
        st.paragraph_format.line_spacing=1
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.left_indent=Pt(0)
        st.paragraph_format.right_indent=Pt(0)
        st.paragraph_format.first_line_indent=Pt(0)
    bibliography=False
    body=False
    for p in list(doc.paragraphs):
        if p.text.strip()=='Introducción':body=True
        if p.text.strip()=='Bibliografía':bibliography=True
        if bibliography:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        elif body and p.style.style_id=='Normal' and p.text.strip():
            p.paragraph_format.space_after=Pt(16.5)
            p.paragraph_format.widow_control=True
        if p.style.style_id==CAP and p.text.startswith('Figura '):
            following=p._p.getnext()
            source=OxmlElement('w:p');p._p.addnext(source)
            q=Paragraph(source,p._parent);q.add_run('Fuente: elaboración propia a partir del diseño y los resultados del proyecto.')
            q.alignment=WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.line_spacing=1;q.paragraph_format.space_after=Pt(8)
            p.paragraph_format.keep_with_next=True
            for r in q.runs:r.font.name='Arial';r.font.size=Pt(10)
    # Each table fits on a page. Keep its rows together, avoiding unlabeled continuations.
    for t in doc.tables:
        if len(t.columns)==1:continue
        for ri,row in enumerate(t.rows):
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next=ri<len(t.rows)-1
        following=t._tbl.getnext()
        if following is not None:
            source=OxmlElement('w:p');t._tbl.addnext(source)
            q=Paragraph(source,doc._body);q.add_run('Fuente: elaboración propia con datos y especificaciones del proyecto.')
            q.alignment=WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.line_spacing=1;q.paragraph_format.space_after=Pt(8)
            for r in q.runs:r.font.name='Arial';r.font.size=Pt(10)
